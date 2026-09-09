"""Build the current boundary-prediction proposal and student handbook."""
from pathlib import Path
import shutil, hashlib, json
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
NAME='低资源条件下句法与构式知识辅助的上海话变调域预测研究.docx'
HANDBOOK_NAME='上海话变调界限预测学习与实验手册.docx'

def configure(d):
    for name,size in [('Normal',11),('Title',20),('Heading 1',15),('Heading 2',12.5)]:
        s=d.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        s.paragraph_format.space_after=Pt(6)
        s.paragraph_format.line_spacing=1.18
        if name!='Normal':s.paragraph_format.keep_with_next=True
    for s in d.sections:
        # Preserve source A4 page and margins. Remove stale footer page fields.
        for p in s.footer.paragraphs:p.clear()
        p=s.footer.paragraphs[0];p.alignment=2
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)

def empty_body(d):
    for e in list(d.element.body):
        if e.tag!=qn('w:sectPr'):d.element.body.remove(e)

def format_cell(cell):
    tcpr=cell._tc.get_or_add_tcPr();borders=tcpr.first_child_found_in('w:tcBorders')
    if borders is None:borders=OxmlElement('w:tcBorders');tcpr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge;node=borders.find(qn(tag))
        if node is None:node=OxmlElement(tag);borders.append(node)
        node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9')

def add_table(d, rows):
    table=d.add_table(rows=1,cols=len(rows[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,text in enumerate(rows[0]):
        cell=table.rows[0].cells[j];cell.text=text;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;format_cell(cell)
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'1F4E78');cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    for i,row in enumerate(rows[1:]):
        cells=table.add_row().cells
        for j,text in enumerate(row):cells[j].text=text;cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;format_cell(cells[j])
        if i%2:
            for cell in cells:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF2F8');cell._tc.get_or_add_tcPr().append(shade)
    d.add_paragraph('')

def add_md(d,path,part=False,part_title=None):
    lines=path.read_text(encoding='utf-8').splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        line=line.strip()
        if not line:continue
        if line.startswith('|') and i<len(lines) and set(lines[i].replace('|','').replace(':','').replace('-','').strip())==set():
            rows=[[c.strip() for c in line.strip('|').split('|')]];i+=1
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]);i+=1
            add_table(d,rows);continue
        if line.startswith('# '):
            if part:d.add_heading(part_title or '变调界限预测学习',level=1)
            else:d.add_paragraph(line[2:],style='Title')
        elif line.startswith('## '):d.add_heading(line[3:],level=1)
        elif line.startswith('### '):d.add_heading(line[4:],level=2)
        else:d.add_paragraph(line)

def main():
    p=ROOT/'docs'/NAME
    if not p.exists():p=ROOT/NAME
    archive=ROOT/'literature/archive'/('原方案修改前_'+NAME)
    archive.parent.mkdir(parents=True,exist_ok=True)
    if not archive.exists():shutil.copy2(p,archive)
    original_hash=hashlib.sha256(archive.read_bytes()).hexdigest()
    d=Document(p);empty_body(d);configure(d);add_md(d,ROOT/'docs/00_原方案修订正文.md');d.save(p)
    teaching=Document(archive);empty_body(teaching);configure(teaching)
    teaching.add_paragraph('上海话变调界限预测学习与实验手册',style='Title')
    add_md(teaching,ROOT/'docs/07_本科生导读与术语例解.md',part=True,part_title='第一部分 本科生导读')
    teaching.add_page_break()
    add_md(teaching,ROOT/'docs/05_学生学习文档.md',part=True,part_title='第二部分 变调界限预测学习')
    teaching.add_page_break()
    add_md(teaching,ROOT/'docs/06_三套理论验证卡.md',part=True,part_title='第三部分 变调界限预测验证卡')
    out=ROOT/'docs'/HANDBOOK_NAME;teaching.save(out)
    (ROOT/'generated/docx_build_manifest.json').write_text(json.dumps({'original_backup':str(archive),'original_sha256':original_hash,'outputs':[str(p),str(out)]},ensure_ascii=False,indent=2),encoding='utf-8')
    print(p);print(out)

if __name__=='__main__':main()
